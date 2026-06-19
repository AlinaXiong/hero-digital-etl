# -*- coding: utf-8 -*-
"""发票信息(DB 直连版)。

按规则表「发票信息」生成增值税发票导入数据:
- 输出字段名取规则表 D 列;
- 必输字段取规则表 I 列=Y;
- 源数据取泛微系统发票台账 fnainvoiceledger,保留已被关联/核销的发票。
- 备注取发票台账扩展表 fnainvoiceledgerdtl.remark。
- 当前只取 2026 年报销/关联的数据。

跑法:在项目根执行  python run.py invoice_info_db
"""
import re
import sys
import time
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

if __package__ is None or __package__ == '':
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from etl import common as c


TASK_NAME = 'invoice_info_db'
OUTPUT_DIR = c.OUT_DIR / TASK_NAME
DATE_SUFFIX = c.today_suffix()
DATA_YEAR = '2026'
DATE_FROM = f'{DATA_YEAR}-01-01'
DATE_TO = '2027-01-01'

RULE_SHEET = '发票信息'
OUTPUT_FILE = OUTPUT_DIR / f'发票信息清洗_发票信息_{DATA_YEAR}_{DATE_SUFFIX}.xlsx'
EXCEPTION_FILE = OUTPUT_DIR / f'未匹配清单_发票信息_{DATA_YEAR}_{DATE_SUFFIX}.xlsx'

SOURCE_TABLE = 'fnainvoiceledger'
SOURCE_DTL_TABLE = 'fnainvoiceledgerdtl'

# fnainvoiceledger.status: 0=初始状态(未使用),1=冻结状态(已被关联),2=核销状态(已核销/已报销)。
RETAINED_STATUS = (1, 2)
QUERY_PARAMS = {
    'retained_status': RETAINED_STATUS,
    'date_from': DATE_FROM,
    'date_to': DATE_TO,
}

FIXED_INVOICE_SOURCE = 'MANUAL'
FIXED_CONFIRM_FLAG = 'N'
FIXED_INVOICE_STATUS = '0'
FIXED_ENTRY_ACCOUNT_STATE = 0
FIXED_VERIFY_FLAG = 0
FIXED_IN_OUT_TYPE = 'IN'
FIXED_REFERENCED_STATUS = 'C'


SOURCE_SQL = """
SELECT
    m.id AS `台账ID`,
    m.userid_new AS `发票归属人ID`,
    m.purchaser AS `购买方`,
    m.purchaserTaxNo AS `购买方纳税人识别号`,
    m.invoiceCode AS `发票代码`,
    m.invoiceNumber AS `发票号码`,
    m.billingDate AS `开票日期`,
    m.taxIncludedPrice AS `价税合计`,
    m.tax AS `税额`,
    m.priceWithoutTax AS `不含税金额`,
    m.seller AS `销售方`,
    m.salesTaxNo AS `销售方纳税人识别号`,
    m.invoiceType AS `发票类型ID`
FROM fnainvoiceledger m
WHERE m.status IN %(retained_status)s
  AND m.reimbursementDate >= %(date_from)s
  AND m.reimbursementDate < %(date_to)s
"""

REMARK_SQL = """
SELECT
    mainId AS `台账ID`,
    MAX(NULLIF(TRIM(remark), '')) AS `备注`
FROM fnainvoiceledgerdtl
WHERE remark IS NOT NULL
  AND TRIM(remark) <> ''
GROUP BY mainId
"""

STATS_SQL = """
SELECT
    COUNT(*) AS total_count,
    SUM(CASE WHEN status IN %(retained_status)s THEN 1 ELSE 0 END) AS retained_all_count,
    SUM(CASE
        WHEN status IN %(retained_status)s
         AND reimbursementDate >= %(date_from)s
         AND reimbursementDate < %(date_to)s
        THEN 1 ELSE 0 END) AS kept_count,
    SUM(CASE WHEN status = 0 THEN 1 ELSE 0 END) AS initial_count
FROM fnainvoiceledger
"""


# 泛微发票类型 -> 汉得 HFINS.VAT_INVOICE_TYPE。
# 规则表只显式写了几类编码;其余按汉得值集 meaning 和台账 category 做可审计映射。
INVOICE_TYPE_MAP = {
    1: '04',                                      # 增值税普通发票
    2: '01',                                      # 增值税专用发票
    3: 'GENERAL_MACHINE_INVOICE',                 # 通用机打发票
    5: 'QUOTA_INVOICE',                           # 定额发票
    7: 'TAXI_INVOICE',                            # 出租发票
    8: '21',                                      # 国际小票/非大陆发票
    9: 'TOLL_INVOICE',                            # 过路费发票
    10: 'BUS_TICKET',                             # 客运汽车发票
    11: '15',                                     # 二手车销售发票
    12: '03',                                     # 机动车销售统一发票
    13: '21',                                     # 国际小票/非大陆发票
    14: 'E-FLIGHT_ITINERARY',                     # 航空运输电子客票行程单
    15: '10',                                     # 增值税电子普通发票
    16: '11',                                     # 增值税普通发票(卷票)
    17: '00',                                     # 可报销其他发票
    18: 'TAX_PAYMENT_RECEIPT',                     # 完税证明
    19: '20',                                     # 区块链发票
    20: '14',                                     # 增值税电子普通发票(通行费)
    21: '08',                                     # 增值税电子专用发票
    22: 'SHIP_TICKET',                            # 船票
    23: '00',                                     # 出行行程单,汉得无专门值集项
    24: 'TRAIN_TICKET_REFUND',                    # 火车票退票凭证
    25: '04',                                     # 电子发票(普通发票),规则表给 04
    32: '31',                                     # 全电专票/数电专票
    33: '32',                                     # 全电普票/数电普票
    34: '00',                                     # 票据汇总单
    35: 'MACHINE_PRINTED_INVOICE',                # 通用(电子)发票
    36: '00',                                     # 门诊收费票据
    37: 'NON_TAX_INCOME_INVOICE',                 # 中央非税收入统一票据
    39: 'CUSTOMS_SPECIAL_PAYMENT_CERTIFICATE',    # 海关缴款书
    47: 'E-TRAIN_TICKET',                         # 电子发票(铁路电子客票)
    48: 'E-FLIGHT_ITINERARY',                     # 电子发票(航空运输电子客票行程单)
    49: '85',                                     # 数电纸质发票(增值税专用发票)
    50: '86',                                     # 数电纸质发票(增值税普通发票)
    51: '21',                                     # 非大陆发票/境外票据
    52: 'E-TRAIN_TICKET_REFUND',                  # 电子发票(铁路电子客票退票凭证)
    53: 'ELECTRONIC_VEHICLE',                     # 电子发票(机动车销售统一发票)
    56: '59',                                     # 电子发票(通行费)
}

ISSUE_SOURCE_FIELD_MAP = {
    '发票归属人中台员工工号': '发票归属人ID',
    '发票归属核算主体中台编码': '购买方',
    '发票票面发票代码': '发票代码',
    '发票票面发票号码': '发票号码',
    '发票票面开票日期': '开票日期',
    '发票票面购买方名称': '购买方',
    '发票票面备注': '备注',
    '发票票面购买方纳税人识别号': '购买方纳税人识别号',
    '发票票面含税金额': '价税合计',
    '发票票面含税金额大写（中文大写）': '价税合计',
    '发票票面税额': '税额',
    '发票票面不含税金额': '不含税金额',
    '发票已被使用含税金额': '价税合计',
    '发票票面销售方名称': '销售方',
    '发票票面销售方纳税人识别号': '销售方纳税人识别号',
    '发票类型，SYSCODE：VAT_INVOICE_TYPE': '发票类型ID',
    '发票来源，SYSCODE：VAT_INVOICE_SOURCE': '发票来源ID',
}


def _text(value):
    """把数据库/Excel 单元格值规整成普通字符串,空值统一返回空串。"""
    if pd.isna(value):
        return ''
    text = str(value).strip()
    return '' if text in ('', 'nan', 'None', 'NaT') else text


def read_rule_columns():
    """从规则表读取输出列和必输字段。D列=字段名,I列=是否必填。"""
    rule_df = pd.read_excel(c.RULE_XLSX, sheet_name=RULE_SHEET, header=None, engine='calamine')
    output_columns = []
    required_columns = []
    for _, row in rule_df.iloc[2:].iterrows():
        field = _text(row[3] if len(row) > 3 else '')
        if not field:
            continue
        output_columns.append(field)
        if _text(row[8] if len(row) > 8 else '') == 'Y':
            required_columns.append(field)
    return output_columns, required_columns


def _decimal_value(value):
    """把金额转成保留两位的 Decimal;无法解析时返回 None。"""
    if pd.isna(value):
        return None
    try:
        return Decimal(str(value).strip()).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None


def amount_to_chinese_upper(value):
    """金额转中文大写,用于目标字段「发票票面含税金额大写」。

    规则表要求 amount_zhs 根据价税合计带出,这里不依赖 Excel 公式,
    直接把数值转成类似「壹佰贰拾叁元肆角伍分」的中文大写。
    """
    amount = _decimal_value(value)
    if amount is None:
        return ''
    if amount == 0:
        return '零元整'

    negative = amount < 0
    amount = abs(amount)
    cents = int((amount * 100).to_integral_value(rounding=ROUND_HALF_UP))
    integer = cents // 100
    fraction = cents % 100

    digits = '零壹贰叁肆伍陆柒捌玖'
    units = ['', '拾', '佰', '仟']
    sections = ['', '万', '亿', '兆']

    def section_to_upper(section):
        """把 0-9999 的四位段转成大写,外层再拼万/亿等节权。"""
        result = ''
        zero = False
        for idx in range(4):
            digit = section % 10
            if digit == 0:
                if result:
                    zero = True
            else:
                if zero:
                    result = digits[0] + result
                    zero = False
                result = digits[digit] + units[idx] + result
            section //= 10
        return result

    int_text = ''
    unit_pos = 0
    need_zero = False
    while integer > 0:
        section = integer % 10000
        if section == 0:
            need_zero = bool(int_text)
        else:
            section_text = section_to_upper(section)
            if need_zero:
                int_text = digits[0] + int_text
                need_zero = False
            int_text = section_text + sections[unit_pos] + int_text
            if section < 1000 and integer >= 10000:
                need_zero = True
        integer //= 10000
        unit_pos += 1

    int_text = int_text or digits[0]
    jiao = fraction // 10
    fen = fraction % 10
    if fraction == 0:
        frac_text = '整'
    elif jiao == 0:
        frac_text = f'零{digits[fen]}分'
    elif fen == 0:
        frac_text = f'{digits[jiao]}角'
    else:
        frac_text = f'{digits[jiao]}角{digits[fen]}分'
    return ('负' if negative else '') + int_text + '元' + frac_text


def _tax_key(value):
    """税号匹配用 key:去掉空格/标点并忽略大小写。"""
    return re.sub(r'[^0-9A-Za-z]', '', _text(value)).casefold()


def _format_date(value):
    """日期容错格式化。非法日期不抛错,返回空值进入缺失清单。"""
    text = _text(value)
    if not text:
        return ''
    date_value = pd.to_datetime(text, errors='coerce')
    return '' if pd.isna(date_value) else date_value.strftime('%Y-%m-%d')


def build_accounting_entity_lookup(purchaser_names, purchaser_tax_numbers):
    """批量构建购买方到汉得核算主体编码的映射。

    返回两个字典:
    - name_map: 购买方名称/纳税人名称(归一化后) -> 核算主体编码;
    - tax_map: 购买方税号(清洗后) -> 核算主体编码。

    后续逐行取值时优先用税号命中,因为名称经常有全半角、括号、简称差异;
    税号为空或查不到时再用购买方名称兜底。
    """
    name_keys = c.normalized_name_values(purchaser_names)
    tax_keys = []
    seen_tax = set()
    for value in purchaser_tax_numbers:
        key = _tax_key(value)
        if key and key not in seen_tax:
            seen_tax.add(key)
            tax_keys.append(key)

    where_parts = []
    params = []
    if name_keys:
        name_expr = c.sql_normalized_name('acc_entity_name')
        taxpayer_name_expr = c.sql_normalized_name('taxpayer_name')
        placeholders = c.in_placeholders(name_keys)
        where_parts.append(f'({name_expr} IN ({placeholders}) OR {taxpayer_name_expr} IN ({placeholders}))')
        params.extend(name_keys + name_keys)
    if tax_keys:
        tax_expr = "LOWER(REGEXP_REPLACE(COALESCE(taxpayer_number, ''), '[^0-9A-Za-z]', ''))"
        where_parts.append(f'{tax_expr} IN ({c.in_placeholders(tax_keys)})')
        params.extend(tax_keys)
    if not where_parts:
        return {}, {}

    entity_df = c.query_db(
        'ZT',
        'hfins_base_account',
        'SELECT acc_entity_code, acc_entity_name, taxpayer_name, taxpayer_number '
        'FROM hfac_accounting_entity '
        f'WHERE {" OR ".join(where_parts)}',
        params,
    )
    name_map = {}
    tax_map = {}
    for _, row in entity_df.iterrows():
        code = _text(row['acc_entity_code'])
        if not code:
            continue
        for name in (row['acc_entity_name'], row['taxpayer_name']):
            key = c.normalize_name(name)
            if key and key not in ('nan', 'none') and key not in name_map:
                name_map[key] = code
        tax_number_key = _tax_key(row['taxpayer_number'])
        if tax_number_key and tax_number_key not in tax_map:
            tax_map[tax_number_key] = code
    return name_map, tax_map


def resolve_accounting_entity_value(purchaser_name, purchaser_tax_number, name_map, tax_map):
    """单行解析核算主体编码:购买方税号优先,购买方名称兜底。"""
    tax_code = tax_map.get(_tax_key(purchaser_tax_number))
    if tax_code:
        return tax_code
    return name_map.get(c.normalize_name(purchaser_name), '')


def resolve_invoice_type(value):
    """把泛微发票类型 ID 转成汉得 HFINS.VAT_INVOICE_TYPE 值集编码。"""
    code = c.format_code(value)
    if not code.isdigit():
        return ''
    return INVOICE_TYPE_MAP.get(int(code), '')


def read_invoice_source():
    """读取泛微发票台账主表。

    当前口径:
    - 只保留 status=1/2 的发票,即已被冻结/关联或已核销的发票;
    - status=0 初始状态视为未使用发票,不导入;
    - 只取 reimbursementDate 落在 2026 年的数据;
    - 备注取 fnainvoiceledgerdtl.remark。

    性能口径:
    - 主表只查输出会用到的字段,不再拉 category/kind 等辅助长字段;
    - 明细表没有 mainId 索引,所以不按 ID 分批反复扫表,而是全表聚合一次后本地 merge。
    """
    start = time.perf_counter()
    stats = c.query_db('FW', 'vspn_xtyy', STATS_SQL, QUERY_PARAMS).iloc[0]
    print(f"[发票信息-DB] 源表 {SOURCE_TABLE}: 总计 {int(stats['total_count'] or 0)} 行; "
          f"剔除初始/未使用 {int(stats['initial_count'] or 0)} 行; "
          f"已关联/已核销 {int(stats['retained_all_count'] or 0)} 行; "
          f"{DATE_FROM}至{DATE_TO}前保留 {int(stats['kept_count'] or 0)} 行")
    print(f'[发票信息-DB] 统计查询耗时: {time.perf_counter() - start:.1f}s')

    start = time.perf_counter()
    source_df = c.query_db('FW', 'vspn_xtyy', SOURCE_SQL, QUERY_PARAMS)
    print(f'[发票信息-DB] 主表查询完成: {len(source_df)} 行, 耗时 {time.perf_counter() - start:.1f}s')

    start = time.perf_counter()
    remark_df = c.query_db('FW', 'vspn_xtyy', REMARK_SQL)
    print(f'[发票信息-DB] 明细备注查询完成: {len(remark_df)} 行, 耗时 {time.perf_counter() - start:.1f}s')
    if not remark_df.empty:
        source_df = source_df.merge(remark_df, on='台账ID', how='left')
        source_df['备注'] = source_df['备注'].fillna('')
    else:
        source_df['备注'] = ''
    print('[发票信息-DB] SQL发票台账行数:', len(source_df))
    return source_df


def build_output(source_df, output_columns):
    """把泛微源字段转换成规则表 D 列要求的目标宽表。

    这里集中处理三类逻辑:
    - 跨系统映射:发票归属人 ID -> 员工工号,购买方 -> 核算主体编码;
    - 字段直取/格式化:发票号、日期、金额、销售方/购买方等;
    - 固定值:发票来源、确认标识、状态、进销项类型、引用状态等。
    """
    print('[发票信息-DB] 解析发票归属人员工工号...')
    employee_map = c.build_fw_employee_info_map_for_ids(source_df['发票归属人ID'])
    print('[发票信息-DB] 匹配购买方核算主体...')
    entity_name_map, entity_tax_map = build_accounting_entity_lookup(
        source_df['购买方'], source_df['购买方纳税人识别号'])

    output_df = pd.DataFrame(index=source_df.index)
    # 规则要求“发票归属人中台员工工号”:泛微 userid_new 先到 hrmresource,
    # 项目现有口径再取 hrmjobtitles.JOBTITLENAME 作为员工工号。
    output_df['发票归属人中台员工工号'] = source_df['发票归属人ID'].map(
        lambda value: employee_map.get(c.format_code(value), {}).get('code', ''))
    # 核算主体不直接存在于泛微台账,用发票购买方税号/名称反查汉得核算主体。
    output_df['发票归属核算主体中台编码'] = [
        resolve_accounting_entity_value(name, tax_number, entity_name_map, entity_tax_map)
        for name, tax_number in zip(source_df['购买方'], source_df['购买方纳税人识别号'])
    ]
    output_df['发票票面发票代码'] = source_df['发票代码'].map(_text)
    output_df['发票票面发票号码'] = source_df['发票号码'].map(_text)
    output_df['发票票面开票日期'] = source_df['开票日期'].map(_format_date)
    output_df['发票票面购买方名称'] = source_df['购买方'].map(_text)
    output_df['发票票面备注'] = source_df['备注'].map(_text)
    output_df['发票票面购买方纳税人识别号'] = source_df['购买方纳税人识别号'].map(_text)
    output_df['发票票面含税金额'] = pd.to_numeric(source_df['价税合计'], errors='coerce').map(c.round_amount)
    print('[发票信息-DB] 生成金额大写...')
    output_df['发票票面含税金额大写（中文大写）'] = source_df['价税合计'].map(amount_to_chinese_upper)
    output_df['发票票面税额'] = pd.to_numeric(source_df['税额'], errors='coerce').map(c.round_amount)
    output_df['发票票面不含税金额'] = pd.to_numeric(source_df['不含税金额'], errors='coerce').map(c.round_amount)
    output_df['发票已被使用含税金额'] = output_df['发票票面含税金额']
    output_df['发票票面销售方名称'] = source_df['销售方'].map(_text)
    output_df['发票票面销售方纳税人识别号'] = source_df['销售方纳税人识别号'].map(_text)
    # 泛微 invoiceType 是内部数字 ID,目标表要的是汉得值集编码。
    output_df['发票类型，SYSCODE：VAT_INVOICE_TYPE'] = source_df['发票类型ID'].map(resolve_invoice_type)

    # 以下字段规则表备注为默认值,不从泛微取值。
    output_df['发票来源，SYSCODE：VAT_INVOICE_SOURCE'] = FIXED_INVOICE_SOURCE
    output_df['发票确认标识'] = FIXED_CONFIRM_FLAG
    output_df['发票状态，SYSCODE：VAT_INVOICE_STATUS'] = FIXED_INVOICE_STATUS
    output_df['入账状态'] = FIXED_ENTRY_ACCOUNT_STATE
    output_df['发票验真标志'] = FIXED_VERIFY_FLAG
    output_df['进销项类型（HFINS.IN_OUT_TYPE）'] = FIXED_IN_OUT_TYPE
    output_df['引用状态(HFINS.REFERENCE_STATUS)'] = FIXED_REFERENCED_STATUS

    for column in output_columns:
        if column not in output_df.columns:
            output_df[column] = ''
    return output_df[output_columns]


def report_fill(output_df, columns):
    """把必输字段填充率打印到控制台,方便跑数时快速判断质量。"""
    for column in columns:
        if column not in output_df.columns:
            continue
        filled_count = (output_df[column].astype(str).str.strip() != '').sum()
        print(f'  {column} 填充率: {filled_count}/{len(output_df)} = {filled_count/len(output_df)*100:.2f}%')


def fill_summary(output_df, required_columns):
    """生成缺失汇总页:只列出填充率未达 100% 的必输字段。"""
    total = len(output_df)
    rows = []
    for column in required_columns:
        column_exists = column in output_df.columns
        filled = int((output_df[column].astype(str).str.strip() != '').sum()) if column_exists else 0
        if filled < total:
            rows.append({'必输字段': column, '填充数': filled, '缺失数': total - filled,
                         '总数': total, '填充率': f'{filled / total * 100:.2f}%',
                         '备注': '' if column_exists else '输出表缺少该列'})
    return pd.DataFrame(rows, columns=['必输字段', '填充数', '缺失数', '总数', '填充率', '备注'])


def collect_field_issues(output_df, source_df, required_columns):
    """生成逐字段缺失明细 sheet。

    只对“部分缺失”的必输字段出明细;整列全空或全满的字段不展开。
    每条明细带出发票号码、台账 ID 和对应泛微源字段,便于回源补数。
    """
    sheets = {}
    total = len(output_df)
    for column in required_columns:
        if column not in output_df.columns:
            continue
        blank_mask = output_df[column].astype(str).str.strip() == ''
        missing_count = int(blank_mask.sum())
        if not (0 < missing_count < total):
            continue
        data = {
            '发票号码': output_df.loc[blank_mask, '发票票面发票号码'].astype(str),
            '台账ID': source_df.loc[blank_mask, '台账ID'].astype(str),
        }
        source_field = ISSUE_SOURCE_FIELD_MAP.get(column)
        if source_field and source_field in source_df.columns:
            data[f'泛微原表-{source_field}'] = source_df.loc[blank_mask, source_field].astype(str)
        sheets[f'缺失_{column[:20]}'] = pd.DataFrame(data).drop_duplicates().reset_index(drop=True)
    return sheets


def write_output(output_df, output_file):
    """直接写出 xlsx,字段名就是规则表 D 列。

    全量约 25 万行,这里用 openpyxl write_only 模式节省内存。
    """
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('发票信息')
    ws.append(list(output_df.columns))
    for row in output_df.itertuples(index=False, name=None):
        ws.append(['' if pd.isna(value) else value for value in row])
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    return output_file


def run():
    """任务入口:读规则 -> 查源表 -> 转目标表 -> 写主文件和缺失清单。"""
    output_columns, required_columns = read_rule_columns()
    source_df = read_invoice_source()
    output_df = build_output(source_df, output_columns)
    print('[发票信息-DB] 输出明细行数:', len(output_df))

    print('[发票信息-DB] 必输字段填充率:')
    report_fill(output_df, required_columns)

    write_output(output_df, OUTPUT_FILE)
    print('已写出:', OUTPUT_FILE)

    sheets = {'必输字段未达100%': fill_summary(output_df, required_columns)}
    sheets.update(collect_field_issues(output_df, source_df, required_columns))
    exception_path = c.write_exceptions(EXCEPTION_FILE, sheets)
    if exception_path:
        print('已写出:', exception_path, '| 各清单条数:', {
            sheet_name: len(sheet_df) for sheet_name, sheet_df in sheets.items()
            if len(sheet_df) > 0
        })

    unmapped_types = sorted(
        c.format_code(value) for value in source_df.loc[
            output_df['发票类型，SYSCODE：VAT_INVOICE_TYPE'].astype(str).str.strip() == '',
            '发票类型ID'
        ].dropna().unique()
    )
    if unmapped_types:
        print('[发票信息-DB] 未映射发票类型ID:', ', '.join(unmapped_types))


if __name__ == '__main__':
    run()
